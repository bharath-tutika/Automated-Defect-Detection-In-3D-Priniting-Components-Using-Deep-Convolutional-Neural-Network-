"""
Comprehensive Excel Export Utility for 3D Printing Defect Detection.
Generates multi-sheet analytical workbooks and single-inspection export sheets
with full formatting, styling, KPI cards, and bounding box telemetry.
"""

import os
import sys
import sqlite3
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import DATABASE_PATH, RESULT_REPORTS_DIR

DEFAULT_MASTER_PATH = RESULT_REPORTS_DIR / "3D_Printing_Defect_Inspection_Results_Output.xlsx"


def get_db_connection():
    """Establish and return SQLite database connection with row factory."""
    if not DATABASE_PATH.exists():
        raise FileNotFoundError(f"Database file not found at: {DATABASE_PATH}")
    conn = sqlite3.connect(str(DATABASE_PATH))
    conn.row_factory = sqlite3.Row
    return conn


def _apply_styles():
    """Return styling dictionary for Excel cell formatting."""
    font_family = "Segoe UI"
    return {
        "title": Font(name=font_family, size=13, bold=True, color="FFFFFF"),
        "subtitle": Font(name=font_family, size=9.5, italic=True, color="E2E8F0"),
        "section": Font(name=font_family, size=11, bold=True, color="1A365D"),
        "header": Font(name=font_family, size=9.5, bold=True, color="FFFFFF"),
        "sub_header": Font(name=font_family, size=9.5, bold=True, color="FFFFFF"),
        "body": Font(name=font_family, size=9, color="2D3748"),
        "bold": Font(name=font_family, size=9, bold=True, color="2D3748"),
        "kpi_lbl": Font(name=font_family, size=8.5, bold=True, color="4A5568"),
        "kpi_val": Font(name=font_family, size=15, bold=True, color="1A365D"),
        "good": Font(name=font_family, size=9, bold=True, color="22543D"),
        "defect": Font(name=font_family, size=9, bold=True, color="742A2A"),
        "fill_navy": PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid"),
        "fill_blue": PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid"),
        "fill_teal": PatternFill(start_color="2C7A7B", end_color="2C7A7B", fill_type="solid"),
        "fill_kpi": PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid"),
        "fill_good": PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid"),
        "fill_defect": PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid"),
        "fill_z1": PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
        "fill_z2": PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid"),
        "border_thin": Border(
            left=Side(style='thin', color='CBD5E0'),
            right=Side(style='thin', color='CBD5E0'),
            top=Side(style='thin', color='CBD5E0'),
            bottom=Side(style='thin', color='CBD5E0')
        ),
        "align_center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "align_left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "align_right": Alignment(horizontal="right", vertical="center", wrap_text=True),
    }


def get_export_filename(filters: dict = None, inspection_id: int = None) -> str:
    """Return a descriptive, sanitized export filename based on active filters."""
    if inspection_id:
        return f"Inspection_Report_#{inspection_id}.xlsx"

    if not filters:
        return "3D_Printing_Defect_Inspection_Results_Output.xlsx"

    search = filters.get("search", "").strip()
    itype = filters.get("type", "all").strip().lower()
    status = filters.get("status", "ALL").strip().upper()
    defect = filters.get("defect", "all").strip()

    active_parts = []
    if defect and defect != "all":
        active_parts.append(f"Defect_{defect.replace(' ', '_')}")
    if status and status != "ALL":
        active_parts.append(f"Status_{status}")
    if itype and itype != "all":
        active_parts.append(f"Type_{itype}")
    if search:
        safe_search = "".join(c for c in search if c.isalnum() or c in "_-")[:15]
        active_parts.append(f"Search_{safe_search}")

    if active_parts:
        return f"Inspection_Export_{'_'.join(active_parts)}.xlsx"
    return "3D_Printing_Defect_Inspection_Results_Output.xlsx"


def generate_master_inspection_excel(output_path: Path = None, filters: dict = None) -> Path:
    """
    Generate the complete or filtered multi-sheet analytical Excel workbook.
    """
    conn = get_db_connection()
    cursor = conn.cursor()

    # Build filtered query if requested
    query = """
        SELECT 
            i.id,
            i.inspection_type,
            i.original_filename,
            i.result_filename,
            i.defect_type,
            i.confidence,
            i.status,
            i.timestamp,
            i.processing_time,
            i.source,
            i.notes,
            COUNT(d.id) AS detection_count
        FROM inspections i
        LEFT JOIN detections d ON i.id = d.inspection_id
        WHERE 1=1
    """
    params = []
    filter_labels = []

    if filters:
        search = filters.get("search", "").strip()
        itype = filters.get("type", "all").strip()
        status = filters.get("status", "ALL").strip()
        defect = filters.get("defect", "all").strip()

        if search:
            query += " AND (i.original_filename LIKE ? OR i.result_filename LIKE ? OR i.defect_type LIKE ? OR i.source LIKE ? OR i.notes LIKE ?)"
            s_pat = f"%{search}%"
            params.extend([s_pat, s_pat, s_pat, s_pat, s_pat])
            filter_labels.append(f"Search: '{search}'")

        if itype and itype.lower() != "all":
            query += " AND i.inspection_type = ?"
            params.append(itype.lower())
            filter_labels.append(f"Type: {itype}")

        if status and status.upper() != "ALL":
            query += " AND i.status = ?"
            params.append(status.upper())
            filter_labels.append(f"Status: {status.upper()}")

        if defect and defect.lower() != "all":
            if defect.lower() in ["normal", "good", "normal / good"]:
                query += " AND (i.defect_type LIKE '%Normal%' OR i.defect_type LIKE '%Good%' OR i.defect_type LIKE '%No Defect%')"
            else:
                query += " AND i.defect_type LIKE ?"
                params.append(f"%{defect}%")
            filter_labels.append(f"Defect: {defect}")

    query += " GROUP BY i.id ORDER BY i.id DESC"
    cursor.execute(query, params)
    inspections_data = cursor.fetchall()

    # Determine unique output path
    if output_path is None:
        RESULT_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        fname = get_export_filename(filters)
        output_path = RESULT_REPORTS_DIR / fname
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

    # Fetch corresponding bounding box detections
    if inspections_data:
        insp_ids = [row["id"] for row in inspections_data]
        placeholders = ",".join(["?"] * len(insp_ids))
        det_query = f"""
            SELECT 
                d.id AS detection_id,
                d.inspection_id,
                i.inspection_type,
                d.class_name,
                d.confidence,
                d.x1,
                d.y1,
                d.x2,
                d.y2,
                ROUND(d.x2 - d.x1, 2) AS box_width,
                ROUND(d.y2 - d.y1, 2) AS box_height,
                ROUND((d.x2 - d.x1) * (d.y2 - d.y1), 2) AS box_area,
                d.timestamp
            FROM detections d
            JOIN inspections i ON d.inspection_id = i.id
            WHERE d.inspection_id IN ({placeholders})
            ORDER BY d.id ASC
        """
        cursor.execute(det_query, insp_ids)
        detections_data = cursor.fetchall()
    else:
        detections_data = []

    conn.close()

    # Build openpyxl workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet
    st = _apply_styles()

    # -------------------------------------------------------------
    # SHEET 1: Executive_Summary_KPIs
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Executive_Summary_KPIs")
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G2")
    ws1["A1"] = "3D PRINTING DEFECT DETECTION - SYSTEM INSPECTION TELEMETRY SUMMARY"
    ws1["A1"].font = st["title"]
    ws1["A1"].fill = st["fill_navy"]
    ws1["A1"].alignment = st["align_center"]

    total_inspections = len(inspections_data)
    good_count = sum(1 for row in inspections_data if row["status"] == "GOOD")
    defect_count = sum(1 for row in inspections_data if row["status"] == "DEFECT")
    total_boxes = len(detections_data)

    avg_conf = (sum(row["confidence"] for row in inspections_data if row["confidence"]) / total_inspections) if total_inspections > 0 else 0.0
    avg_time = (sum(row["processing_time"] for row in inspections_data if row["processing_time"]) / total_inspections) if total_inspections > 0 else 0.0
    defect_rate = (defect_count / total_inspections) if total_inspections > 0 else 0.0

    # KPI Card Blocks
    kpis = [
        ("Total Inspections", f"{total_inspections:,}", "A4", "B5", st["fill_kpi"]),
        ("Defects Identified", f"{defect_count:,}", "C4", "C5", st["fill_defect"]),
        ("Normal / Good Prints", f"{good_count:,}", "D4", "D5", st["fill_good"]),
        ("Defect Rate (%)", f"{defect_rate * 100:.2f}%", "E4", "E5", st["fill_kpi"]),
        ("Mean Model Confidence", f"{avg_conf * 100:.2f}%", "F4", "F5", st["fill_kpi"]),
        ("Avg Processing Time", f"{avg_time:.3f} s", "G4", "G5", st["fill_kpi"]),
    ]

    for title, val, c_top, c_bot, fill in kpis:
        c1 = ws1[c_top]
        c1.value = title.upper()
        c1.font = st["kpi_lbl"]
        c1.alignment = st["align_center"]
        c1.fill = fill
        c1.border = st["border_thin"]

        c2 = ws1[c_bot]
        c2.value = val
        c2.font = st["kpi_val"]
        c2.alignment = st["align_center"]
        c2.fill = fill
        c2.border = st["border_thin"]

    # Defect breakdown table in Sheet 1
    ws1["A7"] = "DEFECT TAXONOMY DISTRIBUTION"
    ws1["A7"].font = st["section"]

    defect_headers = ["Defect Taxonomy Class", "Incident Count", "Proportion (%)", "Severity Index", "Operational Countermeasure"]
    for col_idx, h in enumerate(defect_headers, 1):
        cell = ws1.cell(row=8, column=col_idx, value=h)
        cell.font = st["header"]
        cell.fill = st["fill_blue"]
        cell.alignment = st["align_center"]
        cell.border = st["border_thin"]

    class_counts = {}
    for r in inspections_data:
        dtype = r["defect_type"] or "Unknown"
        class_counts[dtype] = class_counts.get(dtype, 0) + 1

    severities = {
        "Normal": ("Low", "None — Nominal operation"),
        "Stringing": ("Medium", "Increase retraction distance & lower nozzle temperature"),
        "Warping": ("High", "Clean bed with IPA, increase bed temp, add brim/raft"),
        "Layer Shift": ("Critical", "Check belt tension, tighten pulleys, reduce travel speed"),
        "Under-Extrusion": ("High", "Check nozzle clog, calibrate extruder E-steps"),
        "Over-Extrusion": ("Medium", "Lower flow multiplier, calibrate filament diameter"),
        "Spaghetti": ("Catastrophic", "Immediate printer abort; recalibrate first-layer Z-offset"),
    }

    cur_row = 9
    for cls_name, count in sorted(class_counts.items(), key=lambda x: x[1], reverse=True):
        prop = (count / total_inspections * 100) if total_inspections > 0 else 0
        sev, act = severities.get(cls_name, ("Medium", "Inspect print quality"))

        row_vals = [cls_name, count, f"{prop:.2f}%", sev, act]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=cur_row, column=col_idx, value=val)
            cell.font = st["bold"] if col_idx <= 2 else st["body"]
            cell.alignment = st["align_center"] if col_idx in [2, 3, 4] else st["align_left"]
            cell.fill = st["fill_z2"] if cur_row % 2 == 0 else st["fill_z1"]
            cell.border = st["border_thin"]
            if col_idx == 4:
                if sev in ["Critical", "Catastrophic"]:
                    cell.fill = st["fill_defect"]
                    cell.font = st["defect"]
                elif sev == "Low":
                    cell.fill = st["fill_good"]
                    cell.font = st["good"]
        cur_row += 1

    # Auto-fit columns ws1
    for col in ws1.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # -------------------------------------------------------------
    # SHEET 2: Master_Inspections_Log
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Master_Inspections_Log")
    ws2.views.sheetView[0].showGridLines = True

    headers_2 = [
        "Inspection ID",
        "Mode / Type",
        "Timestamp (UTC)",
        "Source Reference",
        "Result Annotated File",
        "Status",
        "Defect Classification",
        "Confidence (%)",
        "Processing Latency (s)",
        "Detections Count",
        "Audit Notes",
    ]

    for col_idx, h in enumerate(headers_2, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = st["header"]
        cell.fill = st["fill_navy"]
        cell.alignment = st["align_center"]
        cell.border = st["border_thin"]

    ws2.row_dimensions[1].height = 24

    for r_idx, row in enumerate(inspections_data, 2):
        conf_pct = f"{(row['confidence'] or 0.0) * 100:.2f}%"
        ptime = f"{row['processing_time']:.3f}" if row["processing_time"] else "N/A"
        status_val = row["status"]

        vals = [
            row["id"],
            str(row["inspection_type"]).capitalize(),
            row["timestamp"],
            row["original_filename"] or "N/A",
            row["result_filename"] or "N/A",
            status_val,
            row["defect_type"] or "Unknown",
            conf_pct,
            ptime,
            row["detection_count"],
            row["notes"] or "-",
        ]

        fill = st["fill_z2"] if r_idx % 2 == 0 else st["fill_z1"]

        for c_idx, v in enumerate(vals, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=v)
            cell.font = st["body"]
            cell.alignment = st["align_center"] if c_idx in [1, 2, 3, 6, 8, 9, 10] else st["align_left"]
            cell.fill = fill
            cell.border = st["border_thin"]

            if c_idx == 1:
                cell.font = st["bold"]
            elif c_idx == 6:  # Status column
                if status_val == "GOOD":
                    cell.fill = st["fill_good"]
                    cell.font = st["good"]
                else:
                    cell.fill = st["fill_defect"]
                    cell.font = st["defect"]

    for col in ws2.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # -------------------------------------------------------------
    # SHEET 3: Bounding_Box_Detections
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Bounding_Box_Detections")
    ws3.views.sheetView[0].showGridLines = True

    headers_3 = [
        "Detection ID",
        "Inspection ID",
        "Inspection Type",
        "Detected Class",
        "Detection Confidence (%)",
        "X1 (px)",
        "Y1 (px)",
        "X2 (px)",
        "Y2 (px)",
        "Box Width (px)",
        "Box Height (px)",
        "Box Area (px²)",
        "Timestamp",
    ]

    for col_idx, h in enumerate(headers_3, 1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.font = st["header"]
        cell.fill = st["fill_blue"]
        cell.alignment = st["align_center"]
        cell.border = st["border_thin"]

    ws3.row_dimensions[1].height = 24

    for r_idx, row in enumerate(detections_data, 2):
        conf_pct = f"{(row['confidence'] or 0.0) * 100:.2f}%"
        fill = st["fill_z2"] if r_idx % 2 == 0 else st["fill_z1"]

        vals = [
            row["detection_id"],
            row["inspection_id"],
            str(row["inspection_type"]).capitalize(),
            row["class_name"],
            conf_pct,
            row["x1"],
            row["y1"],
            row["x2"],
            row["y2"],
            row["box_width"],
            row["box_height"],
            row["box_area"],
            row["timestamp"],
        ]

        for c_idx, v in enumerate(vals, 1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=v)
            cell.font = st["body"]
            cell.alignment = st["align_center"] if c_idx in [1, 2, 3, 5, 13] else (st["align_right"] if 6 <= c_idx <= 12 else st["align_left"])
            cell.fill = fill
            cell.border = st["border_thin"]
            if c_idx in [1, 4]:
                cell.font = st["bold"]

    for col in ws3.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 30)

    # -------------------------------------------------------------
    # SHEET 4: Defect_Class_Distribution
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Defect_Class_Distribution")
    ws4.views.sheetView[0].showGridLines = True

    headers_4 = [
        "Defect Category",
        "Total Inspections",
        "Total Bounding Boxes Detected",
        "Mean Detection Confidence (%)",
        "Min Confidence (%)",
        "Max Confidence (%)",
        "Quality Status",
    ]

    for col_idx, h in enumerate(headers_4, 1):
        cell = ws4.cell(row=1, column=col_idx, value=h)
        cell.font = st["header"]
        cell.fill = st["fill_teal"]
        cell.alignment = st["align_center"]
        cell.border = st["border_thin"]

    ws4.row_dimensions[1].height = 24

    # Compute metrics grouped by defect class
    class_det_stats = {}
    for d in detections_data:
        c = d["class_name"]
        if c not in class_det_stats:
            class_det_stats[c] = []
        class_det_stats[c].append(d["confidence"] or 0.0)

    all_classes = ["Normal", "Stringing", "Warping", "Layer Shift", "Under-Extrusion", "Over-Extrusion", "Spaghetti"]
    for c_name in class_counts:
        if c_name not in all_classes:
            all_classes.append(c_name)

    for r_idx, c_name in enumerate(all_classes, 2):
        c_insp = class_counts.get(c_name, 0)
        c_confs = class_det_stats.get(c_name, [])
        c_boxes = len(c_confs)
        mean_c = (sum(c_confs) / len(c_confs) * 100) if c_confs else 0.0
        min_c = (min(c_confs) * 100) if c_confs else 0.0
        max_c = (max(c_confs) * 100) if c_confs else 0.0
        status_str = "PASS" if c_name == "Normal" else "FAIL"

        vals = [
            c_name,
            c_insp,
            c_boxes,
            f"{mean_c:.2f}%",
            f"{min_c:.2f}%",
            f"{max_c:.2f}%",
            status_str,
        ]

        fill = st["fill_z2"] if r_idx % 2 == 0 else st["fill_z1"]
        for c_idx, v in enumerate(vals, 1):
            cell = ws4.cell(row=r_idx, column=c_idx, value=v)
            cell.font = st["bold"] if c_idx == 1 else st["body"]
            cell.alignment = st["align_center"] if c_idx > 1 else st["align_left"]
            cell.fill = fill
            cell.border = st["border_thin"]
            if c_idx == 7:
                if status_str == "PASS":
                    cell.fill = st["fill_good"]
                    cell.font = st["good"]
                else:
                    cell.fill = st["fill_defect"]
                    cell.font = st["defect"]

    for col in ws4.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws4.column_dimensions[col_letter].width = max(max_len + 3, 16)

    # Save to output path
    wb.save(str(output_path))
    return output_path


def generate_single_inspection_excel(inspection_id: int, output_path: Path = None) -> Path:
    """
    Generate a detailed Excel report workbook for a single specific inspection record.
    """
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT 
            id, inspection_type, original_filename, result_filename,
            defect_type, confidence, status, timestamp, processing_time,
            source, notes
        FROM inspections
        WHERE id = ?
    """, (inspection_id,))
    insp = cursor.fetchone()

    if not insp:
        conn.close()
        raise ValueError(f"Inspection record #{inspection_id} not found in database.")

    cursor.execute("""
        SELECT 
            id, class_name, confidence, x1, y1, x2, y2,
            ROUND(x2 - x1, 2) AS box_width,
            ROUND(y2 - y1, 2) AS box_height,
            ROUND((x2 - x1) * (y2 - y1), 2) AS box_area,
            timestamp
        FROM detections
        WHERE inspection_id = ?
        ORDER BY id ASC
    """, (inspection_id,))
    detections = cursor.fetchall()
    conn.close()

    if output_path is None:
        RESULT_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        output_path = RESULT_REPORTS_DIR / f"Inspection_Report_#{inspection_id}.xlsx"
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Inspection_#{inspection_id}_Report"
    ws.views.sheetView[0].showGridLines = True
    st = _apply_styles()

    # Title Banner
    ws.merge_cells("A1:G2")
    ws["A1"] = f"3D PRINTING DEFECT DETECTION — INSPECTION RECORD #{inspection_id}"
    ws["A1"].font = st["title"]
    ws["A1"].fill = st["fill_navy"]
    ws["A1"].alignment = st["align_center"]

    # Overview Details Table
    ws["A4"] = "INSPECTION METADATA & TELEMETRY"
    ws["A4"].font = st["section"]

    meta_rows = [
        ("Inspection ID", str(insp["id"]), "Status", insp["status"]),
        ("Inspection Type", str(insp["inspection_type"]).capitalize(), "Defect Type", insp["defect_type"] or "Unknown"),
        ("Date & Timestamp", str(insp["timestamp"]), "Confidence", f"{(insp['confidence'] or 0.0) * 100:.2f}%"),
        ("Original File", insp["original_filename"] or "Live Capture", "Processing Latency", f"{insp['processing_time']:.3f} s" if insp["processing_time"] else "N/A"),
        ("Annotated Result", insp["result_filename"] or "N/A", "Total Bounding Boxes", str(len(detections))),
        ("Source Telemetry", insp["source"] or "System Upload", "Audit Notes", insp["notes"] or "Standard inspection run"),
    ]

    cur_row = 5
    for lbl1, val1, lbl2, val2 in meta_rows:
        ws.cell(row=cur_row, column=1, value=lbl1).font = st["bold"]
        ws.cell(row=cur_row, column=1).fill = st["fill_kpi"]
        ws.cell(row=cur_row, column=1).border = st["border_thin"]

        c_v1 = ws.cell(row=cur_row, column=2, value=val1)
        c_v1.font = st["body"]
        c_v1.border = st["border_thin"]

        ws.cell(row=cur_row, column=4, value=lbl2).font = st["bold"]
        ws.cell(row=cur_row, column=4).fill = st["fill_kpi"]
        ws.cell(row=cur_row, column=4).border = st["border_thin"]

        c_v2 = ws.cell(row=cur_row, column=5, value=val2)
        c_v2.font = st["body"]
        c_v2.border = st["border_thin"]

        if lbl2 == "Status":
            if val2 == "GOOD":
                c_v2.fill = st["fill_good"]
                c_v2.font = st["good"]
            else:
                c_v2.fill = st["fill_defect"]
                c_v2.font = st["defect"]

        cur_row += 1

    # Bounding Box Detections Table
    cur_row += 2
    ws.cell(row=cur_row, column=1, value="LOCALIZED BOUNDING BOX ANNOTATIONS").font = st["section"]
    cur_row += 1

    det_headers = ["Detection ID", "Class Name", "Confidence (%)", "X1 (px)", "Y1 (px)", "X2 (px)", "Y2 (px)", "Width (px)", "Height (px)", "Area (px²)"]
    for c_idx, h in enumerate(det_headers, 1):
        cell = ws.cell(row=cur_row, column=c_idx, value=h)
        cell.font = st["header"]
        cell.fill = st["fill_blue"]
        cell.alignment = st["align_center"]
        cell.border = st["border_thin"]

    ws.row_dimensions[cur_row].height = 22

    if not detections:
        cur_row += 1
        ws.cell(row=cur_row, column=1, value="No anomalous defect bounding boxes detected (Nominal Print).").font = st["body"]
    else:
        for d in detections:
            cur_row += 1
            conf_str = f"{(d['confidence'] or 0.0) * 100:.2f}%"
            d_vals = [
                d["id"],
                d["class_name"],
                conf_str,
                d["x1"],
                d["y1"],
                d["x2"],
                d["y2"],
                d["box_width"],
                d["box_height"],
                d["box_area"],
            ]
            fill = st["fill_z2"] if cur_row % 2 == 0 else st["fill_z1"]
            for c_idx, val in enumerate(d_vals, 1):
                cell = ws.cell(row=cur_row, column=c_idx, value=val)
                cell.font = st["body"]
                cell.alignment = st["align_center"] if c_idx in [1, 3] else (st["align_right"] if c_idx >= 4 else st["align_left"])
                cell.fill = fill
                cell.border = st["border_thin"]
                if c_idx == 2:
                    cell.font = st["bold"]

    # Autofit columns
    for col in ws.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    wb.save(str(output_path))
    return output_path
